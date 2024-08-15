import openpyxl as xl

wb=xl.load_workbook('transactions.xlsx')
sheet1=wb['Sheet1']
cell=sheet1['A1']
cell=sheet1.cell(1,1)
print(cell.value)

for row in range(2,sheet1.max_row + 1):
 cell= sheet1.cell(row,3)
 corrected_price = cell.value * 0.9
 corrected_price_cell = sheet1.cell(row,4)
 corrected_price_cell.value=corrected_price

wb.save("transactions_updated.xlsx")


